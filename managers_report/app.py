import os
import uuid
import json
import logging
from datetime import datetime
from dotenv import load_dotenv
from flask import (
    Flask, render_template, request, redirect, url_for,
    session, jsonify, flash, get_flashed_messages, send_file
)
#added chnage all things working
import mysql.connector
from mysql.connector import pooling
import json
import io
from openpyxl import Workbook

# ----------------- Config -----------------
load_dotenv()
APP_SECRET = os.getenv("FLASK_SECRET_KEY", "change-me")

MYSQL_HOST = os.getenv("MYSQL_HOST", "127.0.0.1")
MYSQL_PORT = int(os.getenv("MYSQL_PORT", "3306"))
MYSQL_DB   = os.getenv("MYSQL_DB", "tg_staffbot")
MYSQL_USER = os.getenv("MYSQL_USER", "root")
MYSQL_PASS = os.getenv("MYSQL_PASS", "root@303")

# Setup logging
logging.basicConfig(level=logging.WARNING)
logger = logging.getLogger(__name__)

# ----------------- App -----------------
app = Flask(__name__, template_folder="templates", static_folder="static")
app.secret_key = APP_SECRET

# Register fmt_ist as a Jinja2 global function
def fmt_ist(dt_utc_naive):
    """Display timestamp as-is, assuming it's stored in IST."""
    if not dt_utc_naive:
        return "-"
    return dt_utc_naive.strftime("%d %b %Y %I:%M %p IST")
app.jinja_env.globals['fmt_ist'] = fmt_ist

def _fromjson_filter(value):
    """
    Robustly turn a value into a list for templates:
    - If it's already a list/tuple -> list
    - If it's a JSON string -> parse JSON
    - If parse fails or it's a plain CSV string -> split by comma
    - If None/empty -> []
    This handles MySQL JSON columns (stringified or native) safely.
    """
    if value is None:
        return []
    # If already list/tuple (some drivers may return native types)
    if isinstance(value, (list, tuple)):
        return list(value)
    # If dict, show keys/values as "k:v" (rare, but safe)
    if isinstance(value, dict):
        return [f"{k}:{v}" for k, v in value.items()]
    # Try JSON-decode if it's a string
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return []
        try:
            parsed = json.loads(s)
            if isinstance(parsed, list):
                return parsed
            if isinstance(parsed, dict):
                return [f"{k}:{v}" for k, v in parsed.items()]
            # If it's a scalar JSON, fall back to CSV split logic
        except Exception:
            pass
        # Fallback: treat as CSV string
        return [x.strip() for x in s.split(",") if x.strip()]
    # Any other scalar -> string it
    return [str(value)]

# Make the filter available to Jinja
app.jinja_env.filters['fromjson'] = _fromjson_filter

def _manager_user_id():
    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute("SELECT id FROM users WHERE telegram_id = %s LIMIT 1", (session["manager_tg"],))
        row = cur.fetchone()
        return row["id"] if row else None

def _full_name(first_name, last_name, username, tg):
    fn = (first_name or "").strip()
    ln = (last_name or "").strip()
    full = (fn + " " + ln).strip()
    return full or (username or f"tg:{tg}")


# ----------------- DB Pool -----------------
dbconfig = {
    "host": MYSQL_HOST,
    "port": MYSQL_PORT,
    "database": MYSQL_DB,
    "user": MYSQL_USER,
    "password": MYSQL_PASS,
    "autocommit": True,
}
cnxpool = pooling.MySQLConnectionPool(pool_name="mgrpool", pool_size=5, **dbconfig)

def db_conn():
    return cnxpool.get_connection()

# ----------------- Check for session_token column -----------------
def has_session_token_column():
    try:
        with db_conn() as conn, conn.cursor() as cur:
            cur.execute("SHOW COLUMNS FROM manager_logins LIKE 'session_token'")
            return bool(cur.fetchone())
    except Exception as e:
        logger.warning(f"Error checking session_token column: {e}")
        return False

SESSION_TOKEN_ENABLED = has_session_token_column()

# ----------------- Time Helpers -----------------
def today_ist_str():
    return datetime.now().strftime("%Y-%m-%d")

# ----------------- Auth -----------------
def login_required(view):
    from functools import wraps
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("manager_login"):
            return redirect(url_for("login", next=request.path))
        if SESSION_TOKEN_ENABLED:
            try:
                with db_conn() as conn, conn.cursor(dictionary=True) as cur:
                    cur.execute(
                        "SELECT session_token FROM manager_logins WHERE login = %s LIMIT 1",
                        (session["manager_login"],)
                    )
                    row = cur.fetchone()
                    if not row or row["session_token"] != session.get("session_token"):
                        session.clear()
                        flash("Session expired. Please log in again.", "danger")
                        return redirect(url_for("login", next=request.path))
            except Exception as e:
                logger.warning(f"Session validation error: {e}")
        return view(*args, **kwargs)
    return wrapped

@app.route("/", methods=["GET"])
def root():
    return redirect(url_for("dashboard") if session.get("manager_login") else url_for("login"))

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "GET":
        return render_template("login.html")

    login_id = (request.form.get("login") or "").strip()
    password = request.form.get("password") or ""

    if not login_id or not password:
        flash("Please enter both login and password.", "danger")
        return render_template("login.html", login_prefill=login_id)

    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT login, password, telegram_id, is_active "
            + (", session_token" if SESSION_TOKEN_ENABLED else "") +
            " FROM manager_logins WHERE login = %s LIMIT 1",
            (login_id,)
        )
        row = cur.fetchone()

    if not row:
        flash("No account found for this login.", "danger")
        return render_template("login.html", login_prefill=login_id)

    if row["is_active"] != 1:
        flash("Account is inactive.", "danger")
        return render_template("login.html", login_prefill=login_id)

    if row["password"] != password:
        flash("Incorrect password.", "danger")
        return render_template("login.html", login_prefill=login_id)

    new_token = str(uuid.uuid4()) if SESSION_TOKEN_ENABLED else None
    if SESSION_TOKEN_ENABLED:
        try:
            with db_conn() as conn, conn.cursor() as cur:
                cur.execute(
                    "UPDATE manager_logins SET session_token = %s WHERE login = %s",
                    (new_token, login_id)
                )
        except Exception as e:
            logger.warning(f"Failed to update session_token: {e}")

    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT first_name, last_name "
            "FROM users WHERE telegram_id = %s LIMIT 1",
            (row["telegram_id"],)
        )
        user = cur.fetchone()

    manager_name = (
        f"{user['first_name']} {user['last_name']}".strip() 
        if user and user["first_name"] 
        else row["login"]
    )

    session["manager_login"] = row["login"]
    session["manager_tg"] = row["telegram_id"]
    session["manager_name"] = manager_name
    if SESSION_TOKEN_ENABLED:
        session["session_token"] = new_token
    session.permanent = True

    next_url = request.args.get("next") or url_for("dashboard")
    return redirect(next_url)

@app.route("/logout")
def logout():
    if session.get("manager_login") and SESSION_TOKEN_ENABLED:
        try:
            with db_conn() as conn, conn.cursor() as cur:
                cur.execute(
                    "UPDATE manager_logins SET session_token = NULL WHERE login = %s",
                    (session["manager_login"],)
                )
        except Exception as e:
            logger.warning(f"Failed to clear session_token: {e}")
    session.clear()
    return redirect(url_for("login"))

@app.route("/api/flash-messages", methods=["GET"])
@login_required
def flash_messages():
    messages = get_flashed_messages(with_categories=True)
    return jsonify([{"category": category, "message": message} for category, message in messages])

@app.route("/dashboard")
@login_required
def dashboard():
    return render_template("dashboard.html", default_date=today_ist_str())

# ----------------- NEW VIEW REPORT PAGE (Tabbed) -----------------
@app.route("/view-report", methods=["GET"])
@login_required
def view_report_page():
    # employees (under this manager), plus sites and drones for dropdowns
    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT u.telegram_id, u.first_name, u.last_name, u.username "
            "FROM users u "
            "WHERE u.role = 2 AND u.is_active = 1 AND u.manager_id = "
            "(SELECT id FROM users WHERE telegram_id = %s LIMIT 1) "
            "ORDER BY u.first_name, u.last_name, u.id",
            (session["manager_tg"],)
        )
        emps = cur.fetchall()
        employees = [{
            "telegram_id": e["telegram_id"],
            "name": _full_name(e["first_name"], e["last_name"], e["username"], e["telegram_id"])
        } for e in emps]

        cur.execute("SELECT name FROM master_sites WHERE is_active = 1 ORDER BY name")
        sites = cur.fetchall()

        cur.execute("SELECT name FROM master_drones WHERE is_active = 1 ORDER BY name")
        drones = cur.fetchall()

    # This renders your new tabbed UI template (you added this file separately)
    return render_template("view_report.html", employees=employees, sites=sites, drones=drones)

# ----------------- EXISTING EDIT PAGE ENTRY -----------------
@app.route("/edit-report", methods=["GET", "POST"])
@login_required
def edit_report_page():
    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT u.telegram_id, u.first_name, u.last_name, u.username "
            "FROM users u "
            "WHERE u.role = 2 AND u.is_active = 1 AND u.manager_id = "
            "(SELECT id FROM users WHERE telegram_id = %s LIMIT 1) "
            "ORDER BY u.first_name, u.last_name, u.id",
            (session["manager_tg"],)
        )
        emps = cur.fetchall()

    employees = [
        {
            "telegram_id": e["telegram_id"],
            "name": (f"{e['first_name']} {e['last_name']}".strip() or e["username"] or f"tg:{e['telegram_id']}")
        }
        for e in emps
    ]

    selected_date = request.form.get("date") or request.args.get("date") or today_ist_str()
    selected_employee = request.form.get("employee") or ""

    return render_template(
        "edit_report.html",
        employees=employees,
        selected_date=selected_date,
        selected_employee=selected_employee
    )

@app.route("/api/track")
@login_required
def api_track():
    date_str = request.args.get("date") or today_ist_str()
    mgr_tg_id = session["manager_tg"]

    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT u.telegram_id, u.first_name, u.last_name, u.username "
            "FROM users u "
            "WHERE u.role = 2 AND u.is_active = 1 AND u.manager_id = "
            "(SELECT id FROM users WHERE telegram_id = %s LIMIT 1) "
            "ORDER BY u.first_name, u.last_name, u.id",
            (mgr_tg_id,)
        )
        emps = cur.fetchall()

    tg_ids = [e["telegram_id"] for e in emps]
    names = {}
    for e in emps:
        fname = (e["first_name"] or "").strip()
        lname = (e["last_name"] or "").strip()
        names[e["telegram_id"]] = (fname + " " + lname).strip() or (e["username"] or f"tg:{e['telegram_id']}")

    report_by_tg = {}
    if tg_ids:
        placeholders = ",".join(["%s"] * len(tg_ids))
        q = (
            f"SELECT id, employee_telegram_id, created_at "
            f"FROM reports WHERE report_date = %s AND employee_telegram_id IN ({placeholders})"
        )
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            cur.execute(q, [date_str] + tg_ids)
            for r in cur.fetchall():
                report_by_tg[r["employee_telegram_id"]] = {"id": r["id"], "created_at": r["created_at"]}

    rows = []
    for i, tg_id in enumerate(tg_ids, start=1):
        r = report_by_tg.get(tg_id)
        rows.append({
            "sr": i,
            "employee_telegram_id": tg_id,
            "name": names.get(tg_id, f"tg:{tg_id}"),
            "time": fmt_ist(r["created_at"]) if r else "-",
            "status": "Submitted" if r else "Not Submitted",
        })

    return jsonify({"ok": True, "date": date_str, "rows": rows})

@app.route("/api/reports")
@login_required
def api_reports():
    date_str = request.args.get("date") or today_ist_str()
    employee_tg_id = request.args.get("employee") or ""
    mgr_tg_id = session["manager_tg"]

    if not employee_tg_id:
        return jsonify({"ok": True, "reports": []})

    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT 1 FROM users u "
            "WHERE u.telegram_id = %s AND u.role = 2 AND u.is_active = 1 AND u.manager_id = "
            "(SELECT id FROM users WHERE telegram_id = %s LIMIT 1) LIMIT 1",
            (employee_tg_id, mgr_tg_id)
        )
        if not cur.fetchone():
            return jsonify({"ok": True, "reports": []})

        cur.execute(
            "SELECT id, report_date, site_name, drone_name, pilot_name, copilot_name, "
            "base_height_m, created_at, dgps_used_json, dgps_operators_json, "
            "grid_numbers_json, gcp_points_json, total_area_sq_km, total_time_min, remark "
            "FROM reports WHERE report_date = %s AND employee_telegram_id = %s",
            (date_str, employee_tg_id)
        )
        reports = cur.fetchall()

    return jsonify({"ok": True, "reports": reports})

def _get_report(report_id: int):
    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        cur.execute(
            "SELECT r.* "
            "FROM reports r "
            "WHERE r.id = %s",
            (report_id,)
        )
        rep = cur.fetchone()
        if not rep:
            return None, [], [], []
        cur.execute(
            "SELECT id, flight_time_min, area_sq_km, uav_rover_file, drone_base_file_no "
            "FROM report_flights WHERE report_id = %s ORDER BY id",
            (report_id,)
        )
        flights = cur.fetchall()
        cur.execute(
            "SELECT name FROM master_sites WHERE is_active = 1 ORDER BY name"
        )
        sites = cur.fetchall()
        cur.execute(
            "SELECT name FROM master_drones WHERE is_active = 1 ORDER BY name"
        )
        drones = cur.fetchall()
        return rep, flights, sites, drones

@app.route("/report/<int:report_id>", methods=["GET"])
@login_required
def report_detail(report_id):
    rep, flights, _, _ = _get_report(report_id)
    if not rep:
        flash("Report not found.", "danger")
        return redirect(url_for("dashboard"))
    return render_template("report_detail.html", report=rep, flights=flights, readonly=True)

@app.route("/report/<int:report_id>/preview", methods=["GET"])
@login_required
def report_preview(report_id):
    rep, flights, _, _ = _get_report(report_id)
    if not rep:
        flash("Report not found.", "danger")
        return ("", 404)
    # If modal/JS asks for a fragment, return only the inner markup
    if request.args.get("fragment") == "1" or request.headers.get("X-Requested-With") == "fetch":
        return render_template("report_fragment.html", report=rep, flights=flights)
    # otherwise return the full standalone page (what you already had)
    return render_template("report_detail.html", report=rep, flights=flights)


@app.route("/report/<int:report_id>/edit", methods=["GET", "POST"])
@login_required
def report_edit(report_id):
    rep, flights, sites, drones = _get_report(report_id)
    if not rep:
        flash("Report not found.", "danger")
        return redirect(url_for("edit_report_page"))

    if request.method == "GET":
        return render_template("report_edit.html", report=rep, flights=flights, sites=sites, drones=drones)

    # --- Parse form fields ---
    report_date = (request.form.get("report_date") or "").strip()
    site_name = (request.form.get("site_name") or "").strip()
    drone_name = (request.form.get("drone_name") or "").strip()
    pilot = (request.form.get("pilot_name") or "").strip()
    copilot = (request.form.get("copilot_name") or "").strip()
    remark = (request.form.get("remark") or "").strip()

    def parse_list(val):
        return [x.strip() for x in val.split(",") if x.strip()]

    dgps_used = json.dumps(parse_list(request.form.get("dgps_used") or ""))
    dgps_operators = json.dumps(parse_list(request.form.get("dgps_operators") or ""))
    grid_numbers = json.dumps(parse_list(request.form.get("grid_numbers") or ""))
    gcp_points = json.dumps(parse_list(request.form.get("gcp_points") or ""))

    try:
        base_h = float(request.form.get("base_height_m") or 0)
    except Exception:
        base_h = 0

    # --- Validation ---
    errors = []
    if not report_date: errors.append("Report date is required.")
    if not site_name: errors.append("Site name is required.")
    if not drone_name: errors.append("Drone name is required.")
    if not pilot: errors.append("Pilot name is required.")
    if not copilot: errors.append("Copilot name is required.")
    if not remark: errors.append("Remark is required.")
    if base_h <= 0: errors.append("Base height must be > 0.")
    if not json.loads(dgps_used): errors.append("DGPS used is required.")
    if not json.loads(dgps_operators): errors.append("DGPS operators are required.")
    if not json.loads(grid_numbers): errors.append("Grid numbers are required.")
    if not json.loads(gcp_points): errors.append("GCP points are required.")

    # Flights
    flight_ids = request.form.getlist("flight_id[]")
    flight_times = request.form.getlist("flight_time[]")
    flight_areas = request.form.getlist("flight_area[]")
    flight_ubxs = request.form.getlist("flight_ubx[]")
    flight_bases = request.form.getlist("flight_base[]")

    flights_data = []
    for i in range(len(flight_times)):
        try:
            flight_id = flight_ids[i] if i < len(flight_ids) and flight_ids[i] else None
            time = float(flight_times[i]) if flight_times[i].strip() else 0
            area = float(flight_areas[i]) if flight_areas[i].strip() else 0
            ubx = (flight_ubxs[i] or "").strip()
            base = (flight_bases[i] or "").strip()
            if time < 1: errors.append(f"Flight {i+1}: Time must be ≥ 1.")
            if area <= 0: errors.append(f"Flight {i+1}: Area must be > 0.")
            if not ubx: errors.append(f"Flight {i+1}: UBX is required.")
            if not base: errors.append(f"Flight {i+1}: Base file is required.")
            flights_data.append({
                "id": flight_id,
                "flight_time_min": time,
                "area_sq_km": area,
                "uav_rover_file": ubx,
                "drone_base_file_no": base
            })
        except ValueError:
            errors.append(f"Flight {i+1}: Invalid time or area.")

    if errors:
        return jsonify({"ok": False, "message": "; ".join(errors)})

    # --- Save ---
    try:
        with db_conn() as conn, conn.cursor() as cur:
            cur.execute(
                "UPDATE reports SET report_date=%s, site_name=%s, drone_name=%s, pilot_name=%s, copilot_name=%s, "
                "base_height_m=%s, dgps_used_json=%s, dgps_operators_json=%s, grid_numbers_json=%s, "
                "gcp_points_json=%s, remark=%s, total_time_min=%s, total_area_sq_km=%s "
                "WHERE id=%s",
                (
                    report_date, site_name, drone_name, pilot, copilot, base_h,
                    dgps_used, dgps_operators, grid_numbers, gcp_points, remark,
                    sum(f["flight_time_min"] for f in flights_data),
                    sum(f["area_sq_km"] for f in flights_data),
                    report_id
                )
            )
            cur.execute("DELETE FROM report_flights WHERE report_id=%s", (report_id,))
            for f in flights_data:
                cur.execute(
                    "INSERT INTO report_flights (report_id, flight_time_min, area_sq_km, uav_rover_file, drone_base_file_no) "
                    "VALUES (%s, %s, %s, %s, %s)",
                    (report_id, f["flight_time_min"], f["area_sq_km"], f["uav_rover_file"], f["drone_base_file_no"])
                )
    except Exception as e:
        logger.error(f"Failed to update report {report_id}: {e}")
        return jsonify({"ok": False, "message": "Failed to update report due to a server error."})

    return jsonify({"ok": True, "message": "Report updated successfully."})

@app.route("/report/<int:report_id>/delete", methods=["POST"])
@login_required
def report_delete(report_id):
    try:
        with db_conn() as conn, conn.cursor() as cur:
            cur.execute("DELETE FROM reports WHERE id = %s", (report_id,))
    except Exception as e:
        logger.error(f"Failed to delete report {report_id}: {e}")
        return jsonify({"ok": False, "message": "Failed to delete report due to a server error."})
    return jsonify({"ok": True, "message": "Report deleted successfully."})

# ----------------- NEW APIs for View Reports Tabs -----------------

@app.route("/api/view/date", methods=["GET"])
@login_required
def api_view_date():
    mode = (request.args.get("mode") or "single").lower()
    today = datetime.now().strftime("%Y-%m-%d")

    mgr_id = _manager_user_id()
    if not mgr_id:
        return jsonify({"ok": True, "rows": [], "message": "Manager not found"})

    rows = []
    msg = None

    try:
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            if mode == "range":
                dfrom = request.args.get("from")
                dto = request.args.get("to")
                if not dfrom or not dto:
                    return jsonify({"ok": False, "rows": [], "message": "Please select From and To dates"})
                if dto > today:
                    msg = "Future 'To' date selected. No data."
                    return jsonify({"ok": True, "rows": [], "message": msg})

                cur.execute(
                    "SELECT r.id, r.created_at, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.report_date BETWEEN %s AND %s "
                    "ORDER BY r.created_at ASC, r.id ASC",
                    (mgr_id, dfrom, dto)
                )
            else:
                d = request.args.get("date") or today
                if d > today:
                    msg = "Future date selected. No data."
                    return jsonify({"ok": True, "rows": [], "message": msg})

                cur.execute(
                    "SELECT r.id, r.created_at, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.report_date = %s "
                    "ORDER BY r.created_at ASC, r.id ASC",
                    (mgr_id, d)
                )

            data = cur.fetchall() or []
            for i, r in enumerate(data, start=1):
                rows.append({
                    "sr": i,
                    "first_name": r["first_name"] or "",
                    "last_name": r["last_name"] or "",
                    "id": r["id"],
                })
    except Exception as e:
        logger.warning(f"/api/view/date error: {e}")
        return jsonify({"ok": False, "rows": [], "message": "Server error"})

    return jsonify({"ok": True, "rows": rows, "message": msg})

@app.route("/api/view/employee", methods=["GET"])
@login_required
def api_view_employee():
    tg = request.args.get("employee") or ""
    if not tg:
        return jsonify({"ok": True, "rows": [], "message": "Select an employee"})

    mgr_id = _manager_user_id()
    if not mgr_id:
        return jsonify({"ok": True, "rows": [], "message": "Manager not found"})

    rows = []
    try:
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            # ensure this employee belongs to this manager
            cur.execute(
                "SELECT 1 FROM users WHERE telegram_id = %s AND manager_id = %s LIMIT 1",
                (tg, mgr_id)
            )
            if not cur.fetchone():
                return jsonify({"ok": True, "rows": [], "message": "Employee not under this manager"})

            cur.execute(
                "SELECT id, report_date, site_name, created_at "
                "FROM reports WHERE employee_telegram_id = %s "
                "ORDER BY report_date DESC, created_at DESC, id DESC",
                (tg,)
            )
            data = cur.fetchall() or []
            for i, r in enumerate(data, start=1):
                rows.append({
                    "sr": i,
                    "date": r["report_date"].strftime("%Y-%m-%d") if hasattr(r["report_date"], "strftime") else str(r["report_date"]),
                    "site_name": r["site_name"],
                    "created_at": fmt_ist(r["created_at"]),
                    "id": r["id"]
                })
    except Exception as e:
        logger.warning(f"/api/view/employee error: {e}")
        return jsonify({"ok": False, "rows": [], "message": "Server error"})

    return jsonify({"ok": True, "rows": rows})

@app.route("/api/view/sites", methods=["GET"])
@login_required
def api_view_sites():
    site = (request.args.get("site") or "").strip()
    date_opt = (request.args.get("date") or "").strip()
    if not site:
        return jsonify({"ok": True, "rows": [], "total_area": "0.000", "message": "Select a site"})

    mgr_id = _manager_user_id()
    if not mgr_id:
        return jsonify({"ok": True, "rows": [], "total_area": "0.000", "message": "Manager not found"})

    rows = []
    total_area = 0.0

    try:
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            # list rows
            if date_opt:
                cur.execute(
                    "SELECT r.id, r.report_date, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.site_name = %s AND r.report_date = %s "
                    "ORDER BY r.report_date DESC, r.id DESC",
                    (mgr_id, site, date_opt)
                )
            else:
                cur.execute(
                    "SELECT r.id, r.report_date, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.site_name = %s "
                    "ORDER BY r.report_date DESC, r.id DESC",
                    (mgr_id, site)
                )
            data = cur.fetchall() or []
            for i, r in enumerate(data, start=1):
                rows.append({
                    "sr": i,
                    "first_name": r["first_name"] or "",
                    "last_name": r["last_name"] or "",
                    "date": r["report_date"].strftime("%Y-%m-%d") if hasattr(r["report_date"], "strftime") else str(r["report_date"]),
                    "id": r["id"]
                })

            # total area from report_flights
            if date_opt:
                cur.execute(
                    "SELECT COALESCE(SUM(rf.area_sq_km),0) AS tot "
                    "FROM report_flights rf "
                    "JOIN reports r ON r.id = rf.report_id "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.site_name = %s AND r.report_date = %s",
                    (mgr_id, site, date_opt)
                )
            else:
                cur.execute(
                    "SELECT COALESCE(SUM(rf.area_sq_km),0) AS tot "
                    "FROM report_flights rf "
                    "JOIN reports r ON r.id = rf.report_id "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.site_name = %s",
                    (mgr_id, site)
                )
            tr = cur.fetchone()
            if tr and tr.get("tot") is not None:
                total_area = float(tr["tot"])
    except Exception as e:
        logger.warning(f"/api/view/sites error: {e}")
        return jsonify({"ok": False, "rows": [], "total_area": "0.000", "message": "Server error"})

    return jsonify({"ok": True, "rows": rows, "total_area": f"{total_area:.3f}"})

@app.route("/api/view/drones", methods=["GET"])
@login_required
def api_view_drones():
    drone = (request.args.get("drone") or "").strip()
    date_opt = (request.args.get("date") or "").strip()
    if not drone:
        return jsonify({"ok": True, "rows": [], "total_flights": 0, "message": "Select a drone"})

    mgr_id = _manager_user_id()
    if not mgr_id:
        return jsonify({"ok": True, "rows": [], "total_flights": 0, "message": "Manager not found"})

    rows = []
    total_flights = 0

    try:
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            # list rows
            if date_opt:
                cur.execute(
                    "SELECT r.id, r.report_date, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.drone_name = %s AND r.report_date = %s "
                    "ORDER BY r.report_date DESC, r.id DESC",
                    (mgr_id, drone, date_opt)
                )
            else:
                cur.execute(
                    "SELECT r.id, r.report_date, u.first_name, u.last_name "
                    "FROM reports r "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.drone_name = %s "
                    "ORDER BY r.report_date DESC, r.id DESC",
                    (mgr_id, drone)
                )
            data = cur.fetchall() or []
            for i, r in enumerate(data, start=1):
                rows.append({
                    "sr": i,
                    "first_name": r["first_name"] or "",
                    "last_name": r["last_name"] or "",
                    "date": r["report_date"].strftime("%Y-%m-%d") if hasattr(r["report_date"], "strftime") else str(r["report_date"]),
                    "id": r["id"]
                })

            # total flights from report_flights
            if date_opt:
                cur.execute(
                    "SELECT COUNT(*) AS c "
                    "FROM report_flights rf "
                    "JOIN reports r ON r.id = rf.report_id "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.drone_name = %s AND r.report_date = %s",
                    (mgr_id, drone, date_opt)
                )
            else:
                cur.execute(
                    "SELECT COUNT(*) AS c "
                    "FROM report_flights rf "
                    "JOIN reports r ON r.id = rf.report_id "
                    "JOIN users u ON u.telegram_id = r.employee_telegram_id "
                    "WHERE u.manager_id = %s AND r.drone_name = %s",
                    (mgr_id, drone)
                )
            tr = cur.fetchone()
            if tr and tr.get("c") is not None:
                total_flights = int(tr["c"])
    except Exception as e:
        logger.warning(f"/api/view/drones error: {e}")
        return jsonify({"ok": False, "rows": [], "total_flights": 0, "message": "Server error"})

    return jsonify({"ok": True, "rows": rows, "total_flights": total_flights})

#download working code excel
@app.get("/view-reports/download")
@login_required
def view_reports_download():
    """
    Create an Excel workbook for the active View Reports tab:
      - tab = date | employee | sites | drones
      - Returns an .xlsx with two sheets: 'Summary' and 'Detailed Info'
    Filters (per tab) come in as querystring params.
    """
    tab = (request.args.get("tab") or "").strip().lower()
    if tab not in ("date", "employee", "sites", "drones"):
        return jsonify({"ok": False, "message": "Invalid tab."}), 400

    mgr_tg_id = session["manager_tg"]
    today = datetime.now().strftime("%Y-%m-%d")

    # Helper: get manager's employee tg list + names
    def _team_employees():
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            cur.execute(
                "SELECT u.telegram_id, u.first_name, u.last_name, u.username "
                "FROM users u "
                "WHERE u.role = 2 AND u.is_active = 1 AND u.manager_id = "
                "(SELECT id FROM users WHERE telegram_id = %s LIMIT 1) "
                "ORDER BY u.first_name, u.last_name, u.id",
                (mgr_tg_id,)
            )
            rows = cur.fetchall()
        names = {}
        tgs = []
        for e in rows:
            fn = (e["first_name"] or "").strip()
            ln = (e["last_name"] or "").strip()
            full = (fn + " " + ln).strip() or (e["username"] or f"tg:{e['telegram_id']}")
            names[e["telegram_id"]] = (fn, ln, full)
            tgs.append(e["telegram_id"])
        return tgs, names

    emp_tgs, emp_names = _team_employees()
    if not emp_tgs:
        return jsonify({"ok": False, "message": "No employees assigned."}), 400

    def _rows_date():
        mode = (request.args.get("mode") or "single").lower()
        rows = []
        summary_rows = []  # Sr., first_name, last_name, created_at, id
        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            if mode == "range":
                d_from = (request.args.get("from") or "").strip()
                d_to = (request.args.get("to") or "").strip()
                if not d_from or not d_to:
                    return {"message": "Please select From and To dates.", "rows": [], "summary": [], "range": (d_from, d_to)}
                if d_to > today:
                    return {"message": "Future To date selected. No data.", "rows": [], "summary": [], "range": (d_from, d_to)}
                placeholders = ",".join(["%s"] * len(emp_tgs))
                cur.execute(
                    f"SELECT r.id, r.employee_telegram_id, r.report_date, r.created_at "
                    f"FROM reports r WHERE r.report_date BETWEEN %s AND %s AND r.employee_telegram_id IN ({placeholders}) "
                    f"ORDER BY r.created_at DESC",
                    [d_from, d_to] + emp_tgs
                )
                data = cur.fetchall()
                sr = 1
                for r in data:
                    fn, ln, _ = emp_names.get(r["employee_telegram_id"], ("", "", ""))
                    summary_rows.append([sr, fn, ln, fmt_ist(r["created_at"]), r["id"]])
                    sr += 1
                return {"rows": data, "summary": summary_rows, "mode": "range", "range": (d_from, d_to)}
            else:
                d = (request.args.get("date") or "").strip()
                if not d:
                    return {"message": "Please select a date.", "rows": [], "summary": [], "date": d}
                if d > today:
                    return {"message": "Future date selected. No data.", "rows": [], "summary": [], "date": d}
                placeholders = ",".join(["%s"] * len(emp_tgs))
                cur.execute(
                    f"SELECT r.id, r.employee_telegram_id, r.report_date, r.created_at "
                    f"FROM reports r WHERE r.report_date = %s AND r.employee_telegram_id IN ({placeholders}) "
                    f"ORDER BY r.created_at DESC",
                    [d] + emp_tgs
                )
                data = cur.fetchall()
                sr = 1
                for r in data:
                    fn, ln, _ = emp_names.get(r["employee_telegram_id"], ("", "", ""))
                    summary_rows.append([sr, fn, ln, fmt_ist(r["created_at"]), r["id"]])
                    sr += 1
                return {"rows": data, "summary": summary_rows, "mode": "single", "date": d}

    def _rows_employee():
        tg = (request.args.get("employee") or "").strip()
        if not tg:
            return {"message": "Select an employee.", "rows": [], "summary": [], "emp_name": ""}
        try:
            tg = int(tg)
        except:
            return {"message": "Invalid employee.", "rows": [], "summary": [], "emp_name": ""}

        if tg not in emp_names:
            return {"message": "Employee not in your team.", "rows": [], "summary": [], "emp_name": ""}

        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            cur.execute(
                "SELECT id, report_date, site_name, created_at "
                "FROM reports WHERE employee_telegram_id = %s ORDER BY created_at DESC",
                (tg,)
            )
            data = cur.fetchall()
        # Summary rows: Sr., Date, Site, Submitted at, id
        summary = []
        for i, r in enumerate(data, start=1):
            summary.append([i, r["report_date"], r["site_name"], fmt_ist(r["created_at"]), r["id"]])
        return {"rows": data, "summary": summary, "emp_name": " ".join([x for x in emp_names[tg][:2] if x])}

    def _rows_sites():
        site = (request.args.get("site") or "").strip()
        d = (request.args.get("date") or "").strip()
        if not site:
            return {"message": "Select a site.", "rows": [], "summary": [], "site": site, "date": d, "total_area": 0.0}

        placeholders = ",".join(["%s"] * len(emp_tgs))
        params = []
        where = f"r.site_name = %s AND r.employee_telegram_id IN ({placeholders})"
        params.append(site)
        params += emp_tgs
        if d:
            where += " AND r.report_date = %s"
            params.append(d)

        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            cur.execute(
                f"SELECT r.id, r.employee_telegram_id, r.report_date, r.created_at, r.total_area_sq_km "
                f"FROM reports r WHERE {where} ORDER BY r.created_at DESC",
                params
            )
            data = cur.fetchall()
        summary = []
        total_area = 0.0
        for i, r in enumerate(data, start=1):
            fn, ln, _ = emp_names.get(r["employee_telegram_id"], ("", "", ""))
            summary.append([i, fn, ln, r["report_date"], fmt_ist(r["created_at"]), r["id"]])
            try:
                total_area += float(r["total_area_sq_km"] or 0)
            except:
                pass
        return {"rows": data, "summary": summary, "site": site, "date": d, "total_area": round(total_area, 3)}

    def _rows_drones():
        drone = (request.args.get("drone") or "").strip()
        d = (request.args.get("date") or "").strip()
        if not drone:
            return {"message": "Select a drone.", "rows": [], "summary": [], "drone": drone, "date": d, "total_flights": 0, "total_time": 0}

        placeholders = ",".join(["%s"] * len(emp_tgs))
        params = []
        where = f"r.drone_name = %s AND r.employee_telegram_id IN ({placeholders})"
        params.append(drone)
        params += emp_tgs
        if d:
            where += " AND r.report_date = %s"
            params.append(d)

        with db_conn() as conn, conn.cursor(dictionary=True) as cur:
            cur.execute(
                f"SELECT r.id, r.employee_telegram_id, r.report_date, r.created_at "
                f"FROM reports r WHERE {where} ORDER BY r.created_at DESC",
                params
            )
            data = cur.fetchall()

            # flights count + total time from report_flights
            if data:
                ids = [x["id"] for x in data]
                placeholders2 = ",".join(["%s"] * len(ids))
                cur.execute(
                    f"SELECT rf.report_id, COUNT(*) AS c, SUM(rf.flight_time_min) AS t "
                    f"FROM report_flights rf WHERE rf.report_id IN ({placeholders2}) "
                    f"GROUP BY rf.report_id",
                    ids
                )
                agg = {row["report_id"]: (int(row["c"] or 0), int(row["t"] or 0)) for row in cur.fetchall()}
            else:
                agg = {}

        summary = []
        total_flights = 0
        total_time = 0
        for i, r in enumerate(data, start=1):
            fn, ln, _ = emp_names.get(r["employee_telegram_id"], ("", "", ""))
            c, t = agg.get(r["id"], (0, 0))
            total_flights += c
            total_time += t
            summary.append([i, fn, ln, r["report_date"], fmt_ist(r["created_at"]), r["id"]])
        return {"rows": data, "summary": summary, "drone": drone, "date": d, "total_flights": total_flights, "total_time": total_time}

    # Pull rows & meta for the selected tab
    if tab == "date":
        payload = _rows_date()
    elif tab == "employee":
        payload = _rows_employee()
    elif tab == "sites":
        payload = _rows_sites()
    else:
        payload = _rows_drones()

    rows = payload.get("rows", [])
    summary_rows = payload.get("summary", [])
    if not rows:
        # no data — keep client button disabled normally, but guard here too
        return jsonify({"ok": False, "message": "No data to download for current filters."}), 400

    # ---------- Build workbook ----------
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Summary"
    ws2 = wb.create_sheet("Detailed Info")

    # HEADER + META for Summary
    if tab == "date":
        ws1.append(["Mode", (payload.get("mode") or "").capitalize()])
        if payload.get("mode") == "range":
            f, t = payload.get("range") or ("", "")
            ws1.append(["From", f])
            ws1.append(["To", t])
        else:
            ws1.append(["Date", payload.get("date", "")])
        ws1.append([])  # gap
        ws1.append(["Sr. No.", "Employee First Name", "Employee Last Name", "Report Submitted At", "Report ID"])
        for row in summary_rows:
            ws1.append(row)

    elif tab == "employee":
        ws1.append(["Employee", payload.get("emp_name", "")])
        ws1.append([])
        ws1.append(["Sr. No.", "Date", "Site Name", "Submitted At", "Report ID"])
        for row in summary_rows:
            ws1.append(row)

    elif tab == "sites":
        ws1.append(["Site Name", payload.get("site", "")])
        if payload.get("date"):
            ws1.append(["Date Filter", payload.get("date")])
        ws1.append([])
        ws1.append(["Sr. No.", "Employee First Name", "Employee Last Name", "Date", "Submitted At", "Report ID"])
        for row in summary_rows:
            ws1.append(row)
        ws1.append([])
        ws1.append(["Total Area (sq km)", payload.get("total_area", 0.0)])

    else:  # drones
        ws1.append(["Drone", payload.get("drone", "")])
        ws1.append(["Date", payload.get("date", "") or "-"])
        ws1.append([])
        ws1.append(["Sr. No.", "Employee First Name", "Employee Last Name", "Date", "Submitted At", "Report ID"])
        for row in summary_rows:
            ws1.append(row)
        ws1.append([])
        ws1.append(["Total Flights", payload.get("total_flights", 0)])
        ws1.append(["Total Flight Time (min)", payload.get("total_time", 0)])

    # DETAILED INFO — join reports + report_flights (+ users for names when needed)
    ids = [r["id"] for r in rows]
    placeholders = ",".join(["%s"] * len(ids))
    with db_conn() as conn, conn.cursor(dictionary=True) as cur:
        # pull reports
        cur.execute(
            f"SELECT r.id, r.employee_telegram_id, r.report_date, r.site_name, r.drone_name, r.pilot_name, r.copilot_name, "
            f"r.dgps_used_json, r.dgps_operators_json, r.grid_numbers_json, r.gcp_points_json, "
            f"r.base_height_m, r.total_area_sq_km, r.total_time_min, r.remark, r.created_at "
            f"FROM reports r WHERE r.id IN ({placeholders})",
            ids
        )
        rep_map = {row["id"]: row for row in cur.fetchall()}

        # flights
        cur.execute(
            f"SELECT rf.report_id, rf.flight_time_min, rf.area_sq_km, rf.uav_rover_file, rf.drone_base_file_no "
            f"FROM report_flights rf WHERE rf.report_id IN ({placeholders}) ORDER BY rf.report_id, rf.id",
            ids
        )
        fl_by_rep = {}
        for rf in cur.fetchall():
            fl_by_rep.setdefault(rf["report_id"], []).append(rf)

    # headers for ws2 (no Telegram IDs)
    if tab == "employee":
        ws2.append([
            "Report ID","Date","Site Name","Drone","Pilot","Copilot",
            "DGPS Used","DGPS Operators","Grid Numbers","GCP Points",
            "Base Height (m)","Total Area (sq km)","Total Time (min)","Remark","Submitted At",
            "Flight Time (min)","Flight Area (sq km)","UBX","Base File"
        ])
    else:
        ws2.append([
            "Employee First Name","Employee Last Name",
            "Report ID","Date","Site Name","Drone","Pilot","Copilot",
            "DGPS Used","DGPS Operators","Grid Numbers","GCP Points",
            "Base Height (m)","Total Area (sq km)","Total Time (min)","Remark","Submitted At",
            "Flight Time (min)","Flight Area (sq km)","UBX","Base File"
        ])

    # helper to normalize JSON lists to CSV string
    def _csv(v):
        try:
            if isinstance(v, list):
                return ", ".join([str(x) for x in v])
            if isinstance(v, str):
                import json as _json
                try:
                    parsed = _json.loads(v)
                    if isinstance(parsed, list):
                        return ", ".join([str(x) for x in parsed])
                except Exception:
                    pass
                return v
            return str(v) if v is not None else ""
        except Exception:
            return str(v) if v is not None else ""

    # write detailed rows
    for rid in ids:
        r = rep_map.get(rid)
        if not r:
            continue
        # name fields (avoid telegram id)
        fn, ln, _ = emp_names.get(r["employee_telegram_id"], ("", "", ""))
        # JSON stringify nicely
        used = _csv(r["dgps_used_json"])
        ops  = _csv(r["dgps_operators_json"])
        grid = _csv(r["grid_numbers_json"])
        gcp  = _csv(r["gcp_points_json"])

        flights = fl_by_rep.get(rid, []) or [None]  # at least one line per report
        for rf in flights:
            flight_vals = (
                (rf["flight_time_min"] if rf else ""),
                (rf["area_sq_km"] if rf else ""),
                (rf["uav_rover_file"] if rf else ""),
                (rf["drone_base_file_no"] if rf else "")
            )
            common = [
                r["id"], r["report_date"], r["site_name"], r["drone_name"], r["pilot_name"], r["copilot_name"],
                used, ops, grid, gcp,
                r["base_height_m"], r["total_area_sq_km"], r["total_time_min"], r["remark"], fmt_ist(r["created_at"])
            ]
            if tab == "employee":
                ws2.append([*common, *flight_vals])
            else:
                ws2.append([fn, ln, *common, *flight_vals])

    # stream it
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    fname = f"view_reports_{tab}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=fname)



if __name__ == "__main__":
    app.run(host="0.0.0.0", port=9000, debug=True)
